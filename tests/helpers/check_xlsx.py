"""Test helper: open an .xlsx with openpyxl (an independent implementation) and dump it as JSON."""
import json, sys, zipfile, datetime
from openpyxl import load_workbook

path = sys.argv[1]
with zipfile.ZipFile(path) as z:
    bad = z.testzip()          # verifies every CRC-32 in the archive
    names = z.namelist()
wb = load_workbook(path)       # raises if the package is malformed
out = {"zip_bad_member": bad, "parts": names, "sheets": {}}
for ws in wb.worksheets:
    rows = []
    for row in ws.iter_rows():
        vals = []
        for c in row:
            v = c.value
            if isinstance(v, (datetime.datetime, datetime.date)):
                v = v.isoformat()
            vals.append({"v": v, "fmt": c.number_format, "t": c.data_type})
        rows.append(vals)
    out["sheets"][ws.title] = {"rows": rows, "freeze": ws.freeze_panes, "filter": ws.auto_filter.ref}
print(json.dumps(out))
